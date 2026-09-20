import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

wb=openpyxl.Workbook(); ws=wb.active; ws.title="توقعات المدفوعات"
ws.sheet_view.rightToLeft=True; ws.sheet_view.showGridLines=False
ORANGE="E07C16";ORANGE_L="FCE9D2";INK="26262A";GREY="6D6E70";YELLOW="FFF4CC";GREENL="E4F2E7";HEAD="F3EFEA";WHITE="FFFFFF";BLUEL="E7F0FA"
thin=Side(style="thin",color="D9D7D3");border=Border(left=thin,right=thin,top=thin,bottom=thin)
def c(r,col,v=None,bold=False,size=11,color=INK,fill=None,align="right",fmt=None,bd=True):
    cc=ws.cell(row=r,column=col,value=v);cc.font=Font(name="Arial",bold=bold,size=size,color=color)
    cc.alignment=Alignment(horizontal=align,vertical="center")
    if fill:cc.fill=PatternFill("solid",fgColor=fill)
    if fmt:cc.number_format=fmt
    if bd:cc.border=border
    return cc
for col,w in {"A":3,"B":34,"C":16,"D":16,"E":16,"F":16}.items(): ws.column_dimensions[col].width=w
SAR='#,##0 "﷼"';NUM='#,##0';PCT='0%'

ws.merge_cells("B2:F2");c(2,2,"درب — توقعات مبيعات التطبيق على كل المحطات (لبوابة الدفع)",bold=True,size=15,color=ORANGE,align="right",bd=False)
ws.merge_cells("B3:F3");c(3,2,"مبنية على مبيعات المحطة الفعلية (تقرير Q2 2026) · افتراضات خبير معتمدة",size=10,color=GREY,align="right",bd=False)

gt=5
c(gt,2,"سيناريوهات النمو (3 سنوات)",bold=True,size=12,color=WHITE,fill=ORANGE,bd=False);ws.merge_cells(start_row=gt,start_column=2,end_row=gt,end_column=6)
c(gt+1,4,"السنة 1",bold=True,fill=HEAD,align="center");c(gt+1,5,"السنة 2",bold=True,fill=HEAD,align="center");c(gt+1,6,"السنة 3",bold=True,fill=HEAD,align="center")
c(gt+1,2,"المُدخلات",bold=True,fill=HEAD,align="center");ws.merge_cells(start_row=gt+1,start_column=2,end_row=gt+1,end_column=3)
def yrow(r,label,vals,fmt):
    c(r,2,label,size=11);ws.merge_cells(start_row=r,start_column=2,end_row=r,end_column=3)
    for j,v in enumerate(vals): c(r,4+j,v,fill=YELLOW,align="center",fmt=fmt)
yrow(gt+2,"عدد المحطات",[200,350,500],NUM)
yrow(gt+3,"متوسط مبيعات المحطة / شهر (﷼)",[1000000,850000,750000],SAR)
yrow(gt+4,"نسبة المبيعات عبر التطبيق",[0.30,0.45,0.60],PCT)
ST=gt+2; MON=gt+3; PEN=gt+4
def crow(r,label,f_by_col,fmt,fill):
    c(r,2,label,size=11,bold=True,fill=fill);ws.merge_cells(start_row=r,start_column=2,end_row=r,end_column=3)
    for j,col in enumerate(("D","E","F")): c(r,4+j,f_by_col(col),align="center",fmt=fmt,fill=fill,bold=True)
crow(gt+5,"مبيعات المحطات السنوية (﷼)", lambda col:f"={col}{ST}*{col}{MON}*12", SAR, HEAD)
APPr=gt+6
crow(APPr,"مبيعات التطبيق السنوية (﷼)", lambda col:f"={col}{ST}*{col}{MON}*12*{col}{PEN}", SAR, ORANGE_L)
crow(gt+7,"— منها مدى (﷼)", lambda col:f"={col}{APPr}*0.75", SAR, GREENL)
crow(gt+8,"— منها فيزا (﷼)", lambda col:f"={col}{APPr}*0.25", SAR, GREENL)
crow(gt+9,"عدد عمليات مدى", lambda col:f"=({col}{APPr}*0.75)/120", NUM, BLUEL)
crow(gt+10,"عدد عمليات فيزا", lambda col:f"=({col}{APPr}*0.25)/180", NUM, BLUEL)
crow(gt+11,"إجمالي العمليات", lambda col:f"={col}{gt+9}+{col}{gt+10}", NUM, BLUEL)

c(gt+13,2,"الأصفر = مُدخلات قابلة للتعديل. الافتراضات الثابتة: مدى 75% من القيمة · متوسط تذكرة مدى 120 ﷼ / فيزا 180 ﷼ (مدمجة في المعادلات). المصدر: تقرير Q2 2026.",size=9,color=GREY,align="right",bd=False)
ws.merge_cells(start_row=gt+13,start_column=2,end_row=gt+13,end_column=6)

path="/tmp/claude-0/-home-user-khadija/05ed6524-f029-51fd-96d9-884d1c665dcf/scratchpad/payments_forecast_v3.xlsx"
wb.save(path)
for y,(st,mon,pen) in enumerate([(200,1000000,.30),(350,850000,.45),(500,750000,.60)],1):
    app=st*mon*12*pen
    print(f"Y{y}: app {app:,.0f} | mada {app*.75:,.0f} visa {app*.25:,.0f} | txns {app*.75/120+app*.25/180:,.0f}")
print("saved",path)
